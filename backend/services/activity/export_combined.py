"""เปรียบเทียบผู้เข้าร่วมหลายกิจกรรม (Intersection / Union / ลบ) + Excel export รวม

ภูมิภาคของแผนภาพเวนคำนวณจาก DB ทุกครั้งด้วย `sets.build_regions` — ฝั่ง client ส่งได้แค่
"คีย์ภูมิภาคที่เลือก" กลับมา แล้ว service ตรวจว่าคีย์นั้นมีอยู่จริงก่อนใช้
⇒ เลือกดูคนข้ามห้องผ่านคีย์ปลอมไม่ได้
"""
import io
import math
import time
from datetime import datetime
from typing import Any, Dict, List, Optional, Sequence

import asyncpg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.config import settings
from core.exceptions import ActivityNotFoundError, RoomNotFoundError, ValidationError
from core.privacy import can_view_activity_pii
from core.rbac import require_member, require_permission

from .base import service_logger
from .export import ExportMixin
from .constants import (
    DEFAULT_EXPORT_COL_WIDTH,
    EXPORT_FIELD_WIDTHS,
    EXPORT_HEADER_LABELS,
    PROFILE_FIELDS,
    THAI_TZ,
)
from .sets import (
    MEMBERSHIP_NO,
    MEMBERSHIP_YES,
    build_regions,
    flatten_members,
    regions_by_key,
    select_regions,
    slim_region,
)
from .thai_date import format_buddhist_date

# แผนภาพเวนแบบวงกลมบรรยายได้ครบทุกภูมิภาคเมื่อมี 2–3 วง (2^N − 1 = 3 หรือ 7 ภูมิภาค)
MIN_COMPARE_ACTIVITIES = 2
MAX_COMPARE_ACTIVITIES = 3

# สี/สไตล์ใช้ค่าเดียวกับ export เดี่ยว (export.py) เพื่อให้สองไฟล์หน้าตาเดียวกัน
HEADER_FILL_HEX = "8B5CF6"   # ม่วง (ธีมกิจกรรม)
TOTAL_FILL_HEX = "EDE9FE"    # ม่วงอ่อน (แถวรวม)
ZEBRA_FILL_HEX = "F5F3FF"
BORDER_HEX = "E2E8F0"
MEMBERSHIP_COLOR_HEX = "7C3AED"


def display_nickname(member: dict) -> str:
    """ชื่อเล่นของคนหนึ่ง — ไม่มีก็ถอยไปชื่อจริง แล้วค่อยเลขที่ (ห้ามปล่อยว่าง)"""
    for key in ("nickname", "first_name"):
        value = member.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    fallback = member.get("student_no") or member.get("student_id") or "-"
    return f"#{fallback}"


class ExportCombinedMixin:
    # ------------------------------------------------------------------ helpers

    @staticmethod
    def _validate_compare_ids(activity_ids: Any) -> List[int]:
        """ตรวจ 2–3 กิจกรรม และห้าม id ซ้ำ (คืน id ที่ normalize แล้ว)"""
        if not isinstance(activity_ids, (list, tuple)):
            raise ValidationError("activity_ids ต้องเป็นรายการของรหัสกิจกรรม")
        try:
            ids = [int(a) for a in activity_ids]
        except (TypeError, ValueError) as e:
            raise ValidationError("activity_ids ต้องเป็นจำนวนเต็ม") from e
        if not MIN_COMPARE_ACTIVITIES <= len(ids) <= MAX_COMPARE_ACTIVITIES:
            raise ValidationError(
                f"ต้องเลือก {MIN_COMPARE_ACTIVITIES}–{MAX_COMPARE_ACTIVITIES} กิจกรรม (ส่งมา {len(ids)})"
            )
        if len(set(ids)) != len(ids):
            raise ValidationError("เลือกกิจกรรมซ้ำกัน")
        return ids

    @classmethod
    async def _load_compare_data(
        cls, conn: asyncpg.Connection, room_id: int, activity_ids: Sequence[int]
    ) -> Dict[str, Any]:
        """โหลดกิจกรรม + ผู้เข้าร่วมของทุกกิจกรรมที่เลือก แล้วจัดเป็นภูมิภาค

        ดึงผู้เข้าร่วม **ครั้งเดียว** ด้วย `= ANY($1::int[])` ไม่ใช่ N+1 query
        กิจกรรมที่ไม่อยู่ในห้องนี้ (หรือถูกลบ) ⇒ `ActivityNotFoundError` = กัน IDOR ข้ามห้อง
        """
        activities: List[dict] = []
        for activity_id in activity_ids:
            activity = await cls._fetch_activity_row(conn, room_id, activity_id)
            if not activity:
                raise ActivityNotFoundError(f"ไม่พบกิจกรรม ID: {activity_id}")
            activities.append(activity)

        rows = await conn.fetch(
            f"{cls.PARTICIPANT_SELECT} "
            "WHERE ap.activity_id = ANY($1::int[]) AND ap.deleted_at IS NULL "
            "ORDER BY s.student_no ASC",
            list(activity_ids),
        )
        members_by_activity: Dict[int, List[dict]] = {}
        for row in rows:
            data = dict(row)
            data["metadata"] = cls._parse_metadata(data["metadata"])
            data["earned_hours"] = float(data["earned_hours"] or 0)
            members_by_activity.setdefault(data["activity_id"], []).append(data)

        for activity in activities:
            activity["participant_count"] = len(members_by_activity.get(activity["id"]) or [])

        # 🔴 คีย์ภูมิภาคถูกคำนวณที่นี่เท่านั้น — ค่าที่ client ส่งมาถูก intersect ทีหลัง
        regions = build_regions(list(activity_ids), members_by_activity)
        return {"activities": activities, "regions": regions, "members_by_activity": members_by_activity}

    @classmethod
    def _resolve_selected_regions(cls, regions: List[dict], region_keys: Any) -> List[dict]:
        """ตรวจ `region_keys` กับภูมิภาคที่คำนวณได้จริง แล้วคืนภูมิภาคที่เลือก"""
        wanted = [str(k).strip() for k in (region_keys or []) if str(k).strip()]
        if not wanted:
            raise ValidationError("ต้องเลือกอย่างน้อย 1 ภูมิภาคก่อน export")
        computed = regions_by_key(regions)
        unknown = sorted({k for k in wanted if k not in computed})
        if unknown:
            raise ValidationError(f"ภูมิภาคที่ไม่รู้จัก: {', '.join(unknown)}")
        return select_regions(regions, wanted)

    @staticmethod
    def _build_activity_field_columns(activities: List[dict], only_keys: set) -> List[Dict[str, Any]]:
        """คอลัมน์ฟิลด์เฉพาะกิจกรรม = union ของ required_fields + dynamic_fields ของทุกกิจกรรม

        ป้ายที่คีย์เดียวกันโผล่ในมากกว่า 1 กิจกรรมจะถูกเติมชื่อกิจกรรมนำหน้า — จำเป็นเพราะ
        `df_1` ของกิจกรรม A กับ `df_1` ของกิจกรรม B เป็น **คนละฟิลด์จริง ๆ**
        (และ label อาจซ้ำกันได้แม้คีย์ต่างกัน)
        """
        title_by_id = {activity["id"]: str(activity["title"]) for activity in activities}
        entries: List[Dict[str, Any]] = []
        key_activity_count: Dict[str, int] = {}

        for activity in activities:
            meta = activity.get("metadata") or {}
            dyn_defs = meta.get("dynamic_fields")
            dyn_defs = dyn_defs if isinstance(dyn_defs, list) else []
            dyn_labels: Dict[str, str] = {}
            dyn_types: Dict[str, str] = {}
            for definition in dyn_defs:
                if not isinstance(definition, dict):
                    continue
                key = str(definition.get("key", "")).strip()
                if not key:
                    continue
                dyn_types[key] = str(definition.get("type", "")).strip()
                label = str(definition.get("label", "")).strip()
                if label:
                    dyn_labels[key] = label

            keys: List[str] = []
            required = meta.get("required_fields")
            for key in required if isinstance(required, list) else []:
                key = str(key).strip()
                if key and key not in keys:
                    keys.append(key)
            for key in dyn_labels:
                if key not in keys:
                    keys.append(key)
            if only_keys:
                keys = [key for key in keys if key in only_keys]

            for key in keys:
                entries.append(
                    {
                        "activity_id": activity["id"],
                        "key": key,
                        "label": EXPORT_HEADER_LABELS.get(key) or dyn_labels.get(key) or key,
                        "type": dyn_types.get(key, ""),
                    }
                )
                key_activity_count[key] = key_activity_count.get(key, 0) + 1

        for entry in entries:
            if key_activity_count.get(entry["key"], 0) > 1:
                entry["header"] = f"{title_by_id[entry['activity_id']]} · {entry['label']}"
            else:
                entry["header"] = entry["label"]
        return entries

    @classmethod
    def _combined_field_value(cls, person: dict, entry: Dict[str, Any], can_view_type_a: bool) -> Any:
        """ค่าฟิลด์เฉพาะกิจกรรมของคนหนึ่ง — คนที่ไม่ได้อยู่ในกิจกรรมนั้นได้ค่าว่าง

        🔴 ต้องกันด้วย `activity_ids` ของภูมิภาคก่อนอ่าน metadata เสมอ เพราะคีย์ `df_1`
        มีอยู่ในทุกกิจกรรมแต่เป็นคนละความหมาย ⇒ อ่านข้ามกิจกรรมเมื่อไรก็ได้ค่าผิดทันที
        """
        if entry["activity_id"] not in (person.get("activity_ids") or []):
            return ""
        field = entry["key"]
        if field in PROFILE_FIELDS and not can_view_type_a:
            return ""  # 🔒 PII ของคนที่ยังไม่ยืนยันตัวตน — ไม่ dump ลงไฟล์
        # 🔴 อ่านจาก record ของ **กิจกรรมนั้น** ไม่ใช่ record ที่ติดมากับ region (ดู sets.flatten_members)
        record = (person.get("records_by_activity") or {}).get(entry["activity_id"], person)
        value = cls._field_reader(field)(record)
        if value is None or (isinstance(value, str) and not value.strip()):
            return ""
        # 🌟 ฟิลด์ที่ประกาศ type=datetime → '1 ตุลาคม 2569 13:00 น.'
        value = cls._format_typed_value(field, value, {field: entry["type"]})
        return cls._translate_label(field, value)

    # ------------------------------------------------------------------ compare

    @classmethod
    async def compare_activities(
        cls,
        pool: asyncpg.Pool,
        activity_ids: List[int],
        user_id: int,
        client_source: str,
        actor_identifier: str,
        server_id: Optional[int] = None,
        room_id: Optional[int] = None,
    ) -> dict:
        """เทียบผู้เข้าร่วม 2–3 กิจกรรม → ภูมิภาคของแผนภาพเวน (คืนเฉพาะชื่อ ไม่มี PII)

        เปิดด้วย `require_member` (อ่านอย่างเดียว) เหมือน `get_activity` — สมาชิกห้องดูได้ทุกคน
        ⇒ ห้ามคืน Type A profile fields ออกจากที่นี่เลย และห้ามให้ข้อความ error มี URL ภายใน
        """
        normalized = cls._validate_compare_ids(activity_ids)
        async with pool.acquire() as conn:
            async with conn.transaction():
                target_room_id = await cls._resolve_room_id(conn, server_id, room_id)
                await require_member(conn, target_room_id, user_id)
                data = await cls._load_compare_data(conn, target_room_id, normalized)

        return {
            "activities": [
                {
                    "id": activity["id"],
                    "title": activity["title"],
                    "activity_date": activity["activity_date"],
                    "participant_count": activity["participant_count"],
                }
                for activity in data["activities"]
            ],
            # slim_region = allowlist เฉพาะฟิลด์ชื่อ (endpoint นี้สมาชิกห้องทุกคนเรียกได้)
            "regions": [slim_region(region) for region in data["regions"]],
        }

    # ------------------------------------------------------------------- export

    @classmethod
    async def export_activities_combined_excel(
        cls,
        pool: asyncpg.Pool,
        activity_ids: List[int],
        region_keys: List[str],
        metadata_keys: List[str],
        include_activity_fields: bool,
        user_name: str,
        user_id: int,
        client_source: str,
        actor_identifier: str,
        server_id: Optional[int] = None,
        room_id: Optional[int] = None,
    ) -> io.BytesIO:
        """Export ผู้เข้าร่วมของภูมิภาคที่เลือกเป็น .xlsx (ชีต สรุป + รายชื่อรวม)

        `region_keys` คือหัวใจของการเลือกแบบ Intersection / Union / ลบ — ทั้งหมดคือการเลือก
        ชุดภูมิภาคเดียวกัน (ดู `frontend/src/utils/activitySets.ts` ฝั่งปุ่มสำเร็จรูป)
        """
        start_time = time.time()
        target_room_id = None
        # init ก่อน try กัน audit fallback เจอ unbound variable (บทเรียน "AuditLogger Fallback")
        region_keys_used: List[str] = [str(k).strip() for k in (region_keys or []) if str(k).strip()]
        entity_id = ",".join(str(a) for a in (activity_ids or []))
        try:
            async with pool.acquire() as conn:
                async with conn.transaction():
                    target_room_id = await cls._resolve_room_id(conn, server_id, room_id)
                    await require_permission(conn, target_room_id, user_id, "MANAGE_ACTIVITIES")

                    normalized = cls._validate_compare_ids(activity_ids)
                    data = await cls._load_compare_data(conn, target_room_id, normalized)
                    activities: List[dict] = data["activities"]
                    selected = cls._resolve_selected_regions(data["regions"], region_keys)
                    region_keys_used = [region["key"] for region in selected]
                    people = flatten_members(selected, data["members_by_activity"])
                    if not people:
                        raise ValidationError("ไม่มีผู้เข้าร่วมในภูมิภาคที่เลือก")

                    only_keys = {str(k).strip() for k in (metadata_keys or []) if str(k).strip()}
                    per_activity_fields = (
                        cls._build_activity_field_columns(activities, only_keys)
                        if include_activity_fields
                        else []
                    )
                    has_custom_fields = any(
                        cls._format_custom_fields((person.get("metadata") or {}).get("custom_fields"))
                        for person in people
                    )
                    room = await conn.fetchrow("SELECT room_name FROM rooms WHERE id = $1", target_room_id)
                    room_name = room["room_name"] if room else f"ห้อง #{target_room_id}"
                    title_by_id = {activity["id"]: str(activity["title"]) for activity in activities}

                    wb = Workbook()
                    cls._write_summary_sheet(
                        wb=wb,
                        room_name=room_name,
                        activities=activities,
                        selected=selected,
                        people=people,
                        title_by_id=title_by_id,
                        user_name=user_name,
                    )
                    cls._write_roster_sheet(
                        wb=wb,
                        activities=activities,
                        people=people,
                        per_activity_fields=per_activity_fields,
                        has_custom_fields=has_custom_fields,
                        user_id=user_id,
                    )

                    output = io.BytesIO()
                    wb.save(output)
                    output.seek(0)
                    # 🌟 ชื่อไฟล์ตั้งเป็น attribute บน BytesIO (router อ่านผ่าน getattr) — return type เดิม
                    output.filename = cls._combined_filename([a["title"] for a in activities])

                    exec_time = int((time.time() - start_time) * 1000)
                    await service_logger.log(
                        conn=conn, action="EXPORT", actor_identifier=actor_identifier,
                        client_source=client_source, room_id=target_room_id, user_id=user_id,
                        entity_type="ACTIVITY", entity_id=entity_id, status="success",
                        new_values={
                            "activity_ids": normalized,
                            "region_keys": region_keys_used,
                            "include_activity_fields": include_activity_fields,
                            "metadata_keys": metadata_keys,
                            "total_people": len(people),
                        },
                        endpoint_or_command="export_activities_combined_excel",
                        execution_time_ms=exec_time,
                    )
                    return output
        except Exception as e:
            exec_time = int((time.time() - start_time) * 1000)
            safe_room_id = None if isinstance(e, RoomNotFoundError) else target_room_id
            async with pool.acquire() as error_conn:
                await service_logger.log(
                    conn=error_conn, action="EXPORT", actor_identifier=actor_identifier,
                    client_source=client_source, room_id=safe_room_id, user_id=user_id,
                    entity_type="ACTIVITY", entity_id=entity_id, status="failed",
                    error_detail=str(e),
                    new_values={"region_keys": region_keys_used, "include_activity_fields": include_activity_fields},
                    endpoint_or_command="export_activities_combined_excel",
                    execution_time_ms=exec_time,
                )
            raise e

    @staticmethod
    def _combined_filename(titles: Sequence[Any]) -> str:
        """ชื่อไฟล์จากชื่อกิจกรรมทุกตัว — ตัดที่ 60 ตัวอักษรกัน suffix ถูกตัดทิ้งตอน HTTP"""
        joined = "_".join(ExportMixin._sanitize_filename(title) for title in titles)
        if len(joined) > 60:
            joined = joined[:60].rstrip("_")
        return f"{joined}_รายชื่อผู้เข้าร่วม.xlsx"

    # ------------------------------------------------------------------ sheets

    @staticmethod
    def _write_summary_sheet(
        wb: Workbook,
        room_name: str,
        activities: List[dict],
        selected: List[dict],
        people: List[dict],
        title_by_id: Dict[int, str],
        user_name: str,
    ) -> None:
        """ชีต 1 — ภาพรวมการเปรียบเทียบ (กิจกรรมที่เลือก + ภูมิภาคที่เลือกพร้อมชื่อเล่น)"""
        ws = wb.active
        ws.title = "สรุป"
        ws.sheet_view.showGridLines = False
        for column, width in (("A", 30), ("B", 20), ("C", 64)):
            ws.column_dimensions[column].width = width

        header_fill = PatternFill("solid", fgColor=HEADER_FILL_HEX)
        total_fill = PatternFill("solid", fgColor=TOTAL_FILL_HEX)
        white_bold = Font(bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color=BORDER_HEX),
            right=Side(style="thin", color=BORDER_HEX),
            top=Side(style="thin", color=BORDER_HEX),
            bottom=Side(style="thin", color=BORDER_HEX),
        )

        ws.merge_cells("A1:C1")
        ws["A1"] = f"เปรียบเทียบกิจกรรม — {room_name}"
        ws["A1"].font = Font(bold=True, size=18, color="0F172A")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 34

        ws.merge_cells("A2:C2")
        ws["A2"] = f"สร้างเมื่อ {datetime.now(THAI_TZ).strftime('%d/%m/%Y %H:%M')} น. (เวลาไทย)"
        ws["A2"].font = Font(color="64748B", size=10)
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

        def _section(row: int, label: str) -> None:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            cell = ws.cell(row=row, column=1, value=label)
            cell.font = Font(bold=True, size=12, color="FFFFFF")
            cell.alignment = Alignment(vertical="center")
            for col in (1, 2, 3):
                target = ws.cell(row=row, column=col)
                target.fill = header_fill
                target.border = thin_border
            ws.row_dimensions[row].height = 24

        def _table_header(row: int, labels: Sequence[str]) -> None:
            for idx, label in enumerate(labels, start=1):
                cell = ws.cell(row=row, column=idx, value=label)
                cell.font = white_bold
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border
            ws.row_dimensions[row].height = 24

        def _table_row(row: int, values: Sequence[Any], wrap_width: int = 0) -> None:
            for idx, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=idx, value=value)
                cell.border = thin_border
                if idx == wrap_width:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                else:
                    cell.alignment = Alignment(vertical="center")
            if wrap_width:
                longest = len(str(values[wrap_width - 1] or ""))
                ws.row_dimensions[row].height = max(22, math.ceil(longest / 60) * 15)
            else:
                ws.row_dimensions[row].height = 22

        row = 4
        _section(row, "กิจกรรมที่เลือก")
        row += 1
        _table_header(row, ["ชื่อกิจกรรม", "วันที่", "ผู้เข้าร่วม (คน)"])
        row += 1
        for activity in activities:
            _table_row(
                row,
                [
                    activity["title"],
                    format_buddhist_date(activity["activity_date"]),
                    activity["participant_count"],
                ],
            )
            row += 1

        row += 1
        _section(row, "ภูมิภาคที่เลือก")
        row += 1
        _table_header(row, ["กิจกรรมที่ร่วม", "จำนวน (คน)", "ชื่อเล่น"])
        row += 1
        for region in selected:
            names = ", ".join(display_nickname(member) for member in region["members"])
            label = " + ".join(title_by_id.get(aid, str(aid)) for aid in region["activity_ids"])
            _table_row(row, [label, len(region["members"]), names], wrap_width=3)
            row += 1

        _table_row(row, [f"รวมทั้งหมด (ไม่ซ้ำ): {len(people)} คน", "", ""])
        total_cell = ws.cell(row=row, column=1)
        total_cell.font = Font(bold=True, color="0F172A")
        for col in (1, 2, 3):
            ws.cell(row=row, column=col).fill = total_fill
        ws.row_dimensions[row].height = 24

        row += 2
        ws.cell(row=row, column=1, value=f"ผู้ export: {user_name}").font = Font(color="64748B", size=10)

    @classmethod
    def _write_roster_sheet(
        cls,
        wb: Workbook,
        activities: List[dict],
        people: List[dict],
        per_activity_fields: List[Dict[str, Any]],
        has_custom_fields: bool,
        user_id: int,
    ) -> None:
        """ชีต 2 — รายชื่อรวม: ข้อมูลพื้นฐาน + ติ๊กถูกต่อกิจกรรม + (เลือกได้) ฟิลด์เฉพาะกิจกรรม"""
        ws = wb.create_sheet("รายชื่อรวม")
        ws.sheet_view.showGridLines = False

        header_fill = PatternFill("solid", fgColor=HEADER_FILL_HEX)
        zebra_fill = PatternFill("solid", fgColor=ZEBRA_FILL_HEX)
        total_fill = PatternFill("solid", fgColor=TOTAL_FILL_HEX)
        white_bold = Font(bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color=BORDER_HEX),
            right=Side(style="thin", color=BORDER_HEX),
            top=Side(style="thin", color=BORDER_HEX),
            bottom=Side(style="thin", color=BORDER_HEX),
        )

        headers: List[str] = [
            EXPORT_HEADER_LABELS["student_no"],
            EXPORT_HEADER_LABELS["first_name"],
            EXPORT_HEADER_LABELS["last_name"],
            EXPORT_HEADER_LABELS["nickname"],
        ]
        widths: List[int] = [
            EXPORT_FIELD_WIDTHS["student_no"],
            EXPORT_FIELD_WIDTHS["first_name"],
            EXPORT_FIELD_WIDTHS["last_name"],
            EXPORT_FIELD_WIDTHS["nickname"],
        ]
        kinds: List[str] = ["no", "text", "text", "text"]

        for activity in activities:
            title = str(activity["title"])
            headers.append(title)
            widths.append(max(12, min(28, len(title) + 4)))
            kinds.append("membership")
        for entry in per_activity_fields:
            headers.append(entry["header"])
            widths.append(EXPORT_FIELD_WIDTHS.get(entry["key"], DEFAULT_EXPORT_COL_WIDTH))
            kinds.append("text")
        if has_custom_fields:
            headers.append(EXPORT_HEADER_LABELS["custom_fields"])
            widths.append(EXPORT_FIELD_WIDTHS["custom_fields"])
            kinds.append("wrap")

        ws.append(headers)
        for idx, kind in enumerate(kinds, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]
            cell = ws.cell(row=1, column=idx)
            cell.fill = header_fill
            cell.font = white_bold
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(
                bottom=Side(style="medium", color=HEADER_FILL_HEX),
                top=Side(style="thin", color=BORDER_HEX),
                left=Side(style="thin", color=BORDER_HEX),
                right=Side(style="thin", color=BORDER_HEX),
            )
        ws.row_dimensions[1].height = 32

        is_super_admin = bool(
            settings.SUPER_ADMIN_ID and user_id and int(user_id) == int(settings.SUPER_ADMIN_ID)
        )
        for i, person in enumerate(people, start=2):
            # 🛡️ Consent Model: participant ที่ยังไม่ยืนยันตัวตน → Type A (profile PII) เป็น ""
            can_view_type_a = can_view_activity_pii(
                requester_user_id=user_id,
                target_user_id=person.get("user_id"),
                identity_claimed=person.get("identity_claimed"),
                is_super_admin=is_super_admin,
            )
            person_activity_ids = set(person.get("activity_ids") or [])
            row_values: List[Any] = [
                person.get("student_no"),
                person.get("first_name"),
                person.get("last_name"),
                person.get("nickname"),
            ]
            for activity in activities:
                row_values.append(MEMBERSHIP_YES if activity["id"] in person_activity_ids else MEMBERSHIP_NO)
            for entry in per_activity_fields:
                row_values.append(cls._combined_field_value(person, entry, can_view_type_a))
            if has_custom_fields:
                row_values.append(
                    cls._format_custom_fields((person.get("metadata") or {}).get("custom_fields"))
                )
            ws.append(row_values)

            for col_idx, kind in enumerate(kinds, start=1):
                cell = ws.cell(row=i, column=col_idx)
                cell.border = thin_border
                if kind in ("no", "membership"):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif kind == "wrap":
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                else:
                    cell.alignment = Alignment(vertical="center")
                # 🌟 ติ๊กถูกของกิจกรรมที่คนนี้อยู่ → ตัวหนาสีม่วง ให้กวาดตาแล้วเห็นทันที
                if kind == "membership" and cell.value == MEMBERSHIP_YES:
                    cell.font = Font(bold=True, color=MEMBERSHIP_COLOR_HEX)
                if i % 2 == 0:
                    cell.fill = zebra_fill
        ws.freeze_panes = "A2"

        last_data_row = 1 + len(people)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_data_row}"

        total_row = last_data_row + 1
        merge_to = min(len(headers), 6)
        ws.cell(row=total_row, column=1, value=f"รวมทั้งหมด: {len(people)} คน (ไม่ซ้ำ)")
        if merge_to > 1:
            ws.merge_cells(
                start_row=total_row, start_column=1, end_row=total_row, end_column=merge_to
            )
        total_cell = ws.cell(row=total_row, column=1)
        total_cell.font = Font(bold=True, color="0F172A")
        total_cell.fill = total_fill
        total_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[total_row].height = 24
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=total_row, column=col_idx)
            cell.border = Border(
                top=Side(style="medium", color=HEADER_FILL_HEX),
                left=Side(style="thin", color=BORDER_HEX),
                right=Side(style="thin", color=BORDER_HEX),
                bottom=Side(style="thin", color=BORDER_HEX),
            )
            if 1 < col_idx <= merge_to:
                cell.fill = total_fill
