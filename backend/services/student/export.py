"""Excel export รายชื่อนักเรียน (openpyxl)"""
import asyncpg
import io
import json
import time
from datetime import date, datetime
from typing import Any, Dict, FrozenSet, List, Optional
from zoneinfo import ZoneInfo

from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.config import settings
from core.exceptions import RoomNotFoundError, StudentNotFoundError, ForbiddenError, ValidationError
from core.logger import AuditLogger
from core.name_utils import normalize_nfc, normalize_en, identity_pair
from core.privacy import can_view_pii, mask_private_fields, is_real_claimed_account, PRIVATE_STUDENT_FIELDS
from core.rbac import require_permission, require_member
from services.action_service import ActionService

from .constants import (
    STUDENT_PATCHABLE_COLUMNS, GLOBAL_FIELDS, NAME_NFC_FIELDS, LOCAL_FIELDS,
    THAI_TZ, THAI_MONTH_NAMES, EXPORT_HEADER_LABELS, ROLE_LABELS, STATUS_LABELS,
    DEFAULT_COL_WIDTH, EXPORT_COLUMN_WIDTHS, CENTER_FIELDS, WRAP_TEXT_FIELDS,
)
from .base import service_logger


class ExportMixin:
    @staticmethod
    def _format_buddhist_birthday(value: Any) -> str:
        """วันเกิดแบบไทย: '25 กรกฎาคม 2553' (ปี พ.ศ. = ค.ศ. + 543)."""
        if value is None:
            return ""
        if isinstance(value, datetime):
            value = value.date()
        if not isinstance(value, date):
            return str(value)
        return f"{value.day} {THAI_MONTH_NAMES[value.month - 1]} {value.year + 543}"

    @staticmethod
    def _translate_value(field: str, value: Any) -> Any:
        """แปลงค่าตามชนิดคอลัมน์ (birthday → พ.ศ., class_role/status → ไทย, None → '')."""
        if value is None:
            return ""
        if field == "birthday":
            return ExportMixin._format_buddhist_birthday(value)
        if field == "class_role":
            return ROLE_LABELS.get(str(value), value)
        if field == "status":
            return STATUS_LABELS.get(str(value), value)
        return value

    @classmethod
    def _build_student_workbook(
        cls,
        room_name: str,
        room_code: Optional[str],
        fields: List[str],
        rows: List[dict],
        generated_at: Optional[datetime] = None,
    ) -> io.BytesIO:
        """สร้าง Workbook 2 แผ่น: 'สรุป' (ภาพรวมห้อง) + 'รายชื่อ' (คอลัมน์ตามที่ผู้ใช้เลือก)."""
        if generated_at is None:
            generated_at = datetime.now(THAI_TZ)

        active_count = sum(1 for r in rows if r.get("_status") == "active")
        pending_count = sum(1 for r in rows if r.get("_status") == "pending")
        inactive_count = max(0, len(rows) - active_count - pending_count)
        avg_completion = int(sum(r["_completion_percent"] for r in rows) / len(rows))
        full_count = sum(1 for r in rows if r["_completion_percent"] == 100)

        HEADER_FILL = PatternFill("solid", fgColor="1D4ED8")   # น้ำเงินเข้ม (เหมือน finance)
        TOTAL_FILL = PatternFill("solid", fgColor="D1D5DB")    # เทาอ่อน
        SECTION_FILL = PatternFill("solid", fgColor="EFF6FF")  # ฟ้าอ่อน
        white_bold = Font(bold=True, color="FFFFFF")
        title_font = Font(bold=True, size=16, color="0F172A")

        wb = Workbook()

        # ---- Sheet 1: สรุป ----
        ws_summary = wb.active
        ws_summary.title = "สรุป"
        ws_summary.sheet_view.showGridLines = False
        ws_summary.column_dimensions["A"].width = 34
        ws_summary.column_dimensions["B"].width = 26

        ws_summary["A1"] = f"สรุปข้อมูลนักเรียน — {room_name}"
        ws_summary["A1"].font = title_font
        ws_summary["A2"] = f"สร้างเมื่อ {generated_at.strftime('%d/%m/%Y %H:%M')} น. (เวลาไทย)"
        ws_summary["A2"].font = Font(color="64748B", size=10)

        def _summary_section(row: int, label: str):
            ws_summary.cell(row=row, column=1, value=label).font = Font(bold=True, size=12)
            for col in (1, 2):
                ws_summary.cell(row=row, column=col).fill = SECTION_FILL

        def _summary_row(row: int, label: str, value: Any):
            ws_summary.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws_summary.cell(row=row, column=2, value=value)

        _summary_section(4, "ข้อมูลพื้นฐาน")
        _summary_row(5, "ชื่อห้องเรียน", room_name)
        _summary_row(6, "รหัสห้อง", room_code or "—")
        ws_summary.cell(row=7, column=1, value="จำนวนนักเรียน").font = Font(bold=True)
        count_cell = ws_summary.cell(row=7, column=2, value=len(rows))
        count_cell.font = Font(bold=True, size=12)
        count_cell.fill = TOTAL_FILL
        _summary_row(8, "กำลังเรียน (Active)", active_count)
        _summary_row(9, "รออนุมัติ (Pending)", pending_count)
        _summary_row(10, "พ้นสภาพ (Inactive)", inactive_count)

        _summary_section(12, "ความครบถ้วนของข้อมูล")
        _summary_row(13, "ข้อมูลครบถ้วนเฉลี่ย", avg_completion)
        ws_summary.cell(row=13, column=2).number_format = '0"%"'
        _summary_row(14, "มีข้อมูลครบ 100%", full_count)

        # ---- Sheet 2: รายชื่อ ----
        ws_data = wb.create_sheet("รายชื่อ")
        ws_data.sheet_view.showGridLines = False

        ws_data.append([EXPORT_HEADER_LABELS.get(f, f) for f in fields])
        for idx, field in enumerate(fields, start=1):
            ws_data.column_dimensions[get_column_letter(idx)].width = EXPORT_COLUMN_WIDTHS.get(field, DEFAULT_COL_WIDTH)
            cell = ws_data.cell(row=1, column=idx)
            cell.fill = HEADER_FILL
            cell.font = white_bold
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for i, row in enumerate(rows, start=2):
            ws_data.append([row.get(f, "") for f in fields])
            for col_idx, field in enumerate(fields, start=1):
                cell = ws_data.cell(row=i, column=col_idx)
                if field in CENTER_FIELDS:
                    cell.alignment = Alignment(horizontal="center")
                elif field in WRAP_TEXT_FIELDS:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                if i % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F8FAFC")
        ws_data.freeze_panes = "A2"

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    async def export_students_excel(cls, pool, fields: List[str], user_name: str, user_id: int, client_source: str, actor_identifier: str, server_id: Optional[int] = None, room_id: Optional[int] = None):
        start_time = time.time()
        target_room_id = None
        try:
            async with pool.acquire() as conn:
                async with conn.transaction():
                    target_room_id = await cls.resolve_room_id(conn, server_id, room_id)
                    await require_permission(conn, target_room_id, user_id, "EXPORT_STUDENTS")

                    if not fields:
                        raise ValidationError("กรุณาเลือกอย่างน้อย 1 คอลัมน์")

                    rows = await conn.fetch(f"{cls.BASE_STUDENT_SELECT} WHERE s.room_id = $1 AND s.deleted_at IS NULL ORDER BY s.student_no ASC", target_room_id)
                    if not rows: raise StudentNotFoundError("ไม่พบข้อมูลนักเรียนในห้องนี้")

                    room = await conn.fetchrow("SELECT room_name, room_code FROM rooms WHERE id = $1 AND deleted_at IS NULL", target_room_id)
                    room_name = room["room_name"] if room else f"ห้อง #{target_room_id}"
                    room_code = room["room_code"] if room else None

                    # ป้องกันคอลัมน์ซ้ำ (ถ้า frontend ส่งซ้ำ) รักษาลำดับแรกที่เจอ
                    seen = set()
                    fields = [f for f in fields if not (f in seen or seen.add(f))]

                    is_super_admin = settings.SUPER_ADMIN_ID and user_id and int(user_id) == int(settings.SUPER_ADMIN_ID)
                    data = []
                    for r in rows:
                        row_dict = dict(r)
                        processed = {f: cls._translate_value(f, row_dict.get(f)) for f in fields}
                        # คีย์ภายในสำหรับ Sheet สรุป (ไม่ถูกเขียนลง Sheet รายชื่อ)
                        processed["_status"] = row_dict.get("status")
                        processed["_completion_percent"] = cls._calculate_completion(row_dict)["percentage"]
                        # 🛡️ Consent Model: สมาชิกที่ยังไม่ยืนยันตัวตน → PII เป็น blank ในไฟล์ export
                        # (คำนวณ completion จาก row_dict ดิบก่อนแล้ว — ไม่ให้ % เปลี่ยนเพราะ mask)
                        if not (is_super_admin or row_dict.get("identity_claimed")):
                            for f in PRIVATE_STUDENT_FIELDS:
                                if f in processed:
                                    processed[f] = ""
                        data.append(processed)

                    output = cls._build_student_workbook(
                        room_name=room_name,
                        room_code=room_code,
                        fields=fields,
                        rows=data,
                        generated_at=datetime.now(THAI_TZ),
                    )

                    exec_time = int((time.time() - start_time) * 1000)
                    await service_logger.log(
                        conn=conn, action="EXPORT", actor_identifier=actor_identifier,
                        client_source=client_source, room_id=target_room_id, user_id=user_id,
                        entity_type="STUDENT_LIST", status="success", new_values={"fields": fields},
                        endpoint_or_command="export_students_excel", execution_time_ms=exec_time
                    )

                    return output
        except Exception as e:
            exec_time = int((time.time() - start_time) * 1000)
            async with pool.acquire() as error_conn:
                await service_logger.log(
                    conn=error_conn, action="EXPORT", actor_identifier=actor_identifier,
                    client_source=client_source, room_id=target_room_id, user_id=user_id,
                    entity_type="STUDENT_LIST", status="failed", error_detail=str(e),
                    new_values={"fields": fields}, endpoint_or_command="export_students_excel", execution_time_ms=exec_time
                )
            raise e
